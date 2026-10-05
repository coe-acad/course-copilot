"""Convert markdown asset content into a formatted Excel workbook.

Markdown tables become real, bordered spreadsheet tables; headings, paragraphs,
lists and code are laid out as rows above/below them. A sprint-structure
timetable (a table with "Coffee Break" / "Lunch" columns) additionally gets the
activity colour-coding used in the app, merged break columns and grey weekend rows.
"""

from __future__ import annotations

import math
import re
from io import BytesIO
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import sprint_structure as ss
from .text_to_docx import _split_blocks, latex_to_text

FONT_NAME = "Calibri"
BASE_SIZE = 11
MIN_COL_WIDTH = 10
MAX_COL_WIDTH = 48
PROSE_WIDTH = 110  # total width prose rows are wrapped to when there are no tables
LINE_HEIGHT = 15  # points per wrapped line at 11pt

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BLACK_SIDE = Side(style="thin", color="000000")
SPRINT_BORDER = Border(left=BLACK_SIDE, right=BLACK_SIDE, top=BLACK_SIDE, bottom=BLACK_SIDE)
HEADER_FILL = PatternFill("solid", fgColor="F2F2F2")
TOP_WRAP = Alignment(wrap_text=True, vertical="top")

WEEKEND_FILL = PatternFill("solid", fgColor=ss.WEEKEND_BG)

_normalise = ss.normalise
_sprint_colour = ss.activity_colour
_is_break = ss.is_break
_is_sprint_table = ss.is_sprint_table


def _plain(text: str) -> str:
    """Strip inline markdown so cells hold clean text."""
    if not text:
        return ""
    t = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    t = re.sub(r"!\[([^\]]*)\]\([^)]*\)", lambda m: f"[Image: {m.group(1)}]" if m.group(1) else "[Image]", t)
    t = re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r"\1 (\2)", t)
    t = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), t)
    t = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?!\w)", r"\1", t)
    t = re.sub(r"(?<!\w)_(?!\s)(.+?)(?<!\s)_(?!\w)", r"\1", t)
    t = re.sub(r"`([^`]+)`", r"\1", t)
    t = t.replace("\\|", "|")
    return t.strip()


def _is_all_bold(text: str) -> bool:
    return bool(re.fullmatch(r"\s*(\*\*|__)(.+)\1\s*", text or ""))


def _estimate_lines(text: str, width: float) -> int:
    """Wrapped line count of text in a column of the given Excel width."""
    chars_per_line = max(int(width * 1.1), 1)
    return sum(max(1, math.ceil(len(part) / chars_per_line)) for part in (text or "").split("\n"))


class _SheetWriter:
    def __init__(self, ws, total_cols: int, col_widths: List[float]):
        self.ws = ws
        self.row = 1
        self.total_cols = max(total_cols, 1)
        self.col_widths = col_widths

    def _prose_width(self) -> float:
        return sum(self.col_widths[: self.total_cols])

    def _merge_row(self):
        if self.total_cols > 1:
            self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.total_cols)

    def prose(self, text: str, *, bold=False, size=BASE_SIZE, italic=False, color="1F2937",
              fill: Optional[str] = None, mono=False, gap_after=1):
        if not text.strip():
            return
        cell = self.ws.cell(row=self.row, column=1, value=text)
        cell.font = Font(name="Consolas" if mono else FONT_NAME, size=size, bold=bold, italic=italic, color=color)
        cell.alignment = TOP_WRAP
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        self._merge_row()
        lines = _estimate_lines(text, self._prose_width() * BASE_SIZE / size)
        self.ws.row_dimensions[self.row].height = max(lines * LINE_HEIGHT * size / BASE_SIZE, LINE_HEIGHT)
        self.row += 1 + gap_after

    def table(self, headers: List[str], rows: List[List[str]], sprint: bool, sprint_doc: bool):
        ncols = len(headers)
        rows = [(r + [""] * ncols)[:ncols] for r in rows]
        start_row = self.row

        # Header row
        for c, h in enumerate(headers, start=1):
            cell = self.ws.cell(row=self.row, column=c, value=_plain(h))
            if sprint or sprint_doc:
                cell.font = Font(name=FONT_NAME, size=BASE_SIZE, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="000000")
                cell.border = SPRINT_BORDER
                cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center",
                                           text_rotation=90 if sprint and _is_break(_plain(h)) else 0)
            else:
                cell.font = Font(name=FONT_NAME, size=BASE_SIZE, bold=True, color="111827")
                cell.fill = HEADER_FILL
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        # Sprint headers are tall enough for the rotated "Coffee Break" labels.
        self._fit_row_height([_plain(h) for h in headers], minimum=78 if sprint else LINE_HEIGHT)
        self.row += 1
        body_start = self.row

        for r in rows:
            weekend = sprint and _normalise(_plain(r[0])) == "weekend"
            plain_cells = [_plain(v) for v in r]
            for c, raw in enumerate(r, start=1):
                text = plain_cells[c - 1]
                cell = self.ws.cell(row=self.row, column=c)
                if sprint or sprint_doc:
                    self._sprint_cell(cell, raw, text, weekend)
                else:
                    cell.value = text
                    cell.font = Font(name=FONT_NAME, size=BASE_SIZE, bold=_is_all_bold(raw), color="1F2937")
                    cell.border = BORDER
                    cell.alignment = TOP_WRAP
            if weekend:
                self.ws.row_dimensions[self.row].height = 18
            else:
                self._fit_row_height(plain_cells, extra_line=sprint)
            self.row += 1

        if sprint:
            self._merge_break_columns(headers, body_start, self.row - 1)
            self.ws.freeze_panes = self.ws.cell(row=body_start, column=3)
        elif rows:
            self.ws.auto_filter.ref = f"A{start_row}:{get_column_letter(ncols)}{self.row - 1}"
        self.row += 1

    def _sprint_cell(self, cell, raw: str, text: str, weekend: bool):
        cell.border = SPRINT_BORDER
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        if weekend:
            cell.value = text
            cell.fill = WEEKEND_FILL
            cell.font = Font(name=FONT_NAME, size=BASE_SIZE, bold=True, color="4B5563")
            return
        colour = _sprint_colour(text) if text else None
        if colour:
            cell.fill = PatternFill("solid", fgColor=colour[0])
        font_color = colour[1] if colour else "111827"
        bold = bool(colour and colour[2])  # only activities the legend marks bold
        split = re.match(r"^(.+?)\s+[–—-]\s+(.+)$", text, flags=re.S)
        if colour and split and _sprint_colour(split.group(1)):
            # "Label – topic": activity label on its own line above the topic.
            cell.value = CellRichText(
                TextBlock(InlineFont(rFont=FONT_NAME, sz=BASE_SIZE, b=bold, color=font_color), split.group(1)),
                TextBlock(InlineFont(rFont=FONT_NAME, sz=BASE_SIZE - 1, b=bold, color=font_color), "\n" + split.group(2)),
            )
        else:
            cell.value = text
            cell.font = Font(name=FONT_NAME, size=BASE_SIZE, bold=bold, color=font_color)

    def _merge_break_columns(self, headers: List[str], first: int, last: int):
        """Merge each break column vertically within a week (between weekend rows)."""
        for c, h in enumerate(headers, start=1):
            if not _is_break(_plain(h)):
                continue
            run_start = None
            for r in range(first, last + 2):
                is_break_cell = r <= last and _is_break(str(self.ws.cell(row=r, column=c).value or ""))
                if is_break_cell and run_start is None:
                    run_start = r
                elif not is_break_cell and run_start is not None:
                    top = self.ws.cell(row=run_start, column=c)
                    if r - 1 > run_start:
                        self.ws.merge_cells(start_row=run_start, start_column=c, end_row=r - 1, end_column=c)
                    top.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90)
                    top.font = Font(name=FONT_NAME, size=BASE_SIZE, bold=True, color="374151")
                    run_start = None

    def _fit_row_height(self, texts: List[str], minimum: float = LINE_HEIGHT, extra_line=False):
        lines = 1
        for c, text in enumerate(texts):
            if _is_break(text):
                continue  # rotated / merged; don't let it stretch the row
            width = self.col_widths[c] if c < len(self.col_widths) else MIN_COL_WIDTH
            lines = max(lines, _estimate_lines(text, width))
        height = lines * LINE_HEIGHT + (LINE_HEIGHT if extra_line and lines > 1 else 6)
        self.ws.row_dimensions[self.row].height = max(height, minimum)


def _column_widths(blocks: List[dict]) -> List[float]:
    """Column widths sized to the widest table content, capped for readability."""
    widths: List[float] = []
    has_sprint = any(b["type"] == "table" and _is_sprint_table(b["headers"]) for b in blocks)
    for b in blocks:
        if b["type"] != "table":
            continue
        if has_sprint and len(b["headers"]) == 1:
            continue  # the sprint legend wraps inside the Date column
        sprint = _is_sprint_table(b["headers"])
        for row in [b["headers"]] + b["rows"]:
            for c, raw in enumerate(row):
                text = _plain(raw)
                if not text:
                    continue  # blank cells (e.g. weekend rows) shouldn't widen a column
                if sprint and _is_break(text):
                    want = 6.5
                else:
                    longest = max((len(p) for p in text.split("\n")), default=0)
                    if sprint:
                        # Compact timetable: short columns (Day) stay narrow, long cells wrap.
                        want = min(max(longest * 1.1 + 2, 6), 24)
                    else:
                        want = min(max(longest * 1.1 + 2, MIN_COL_WIDTH), MAX_COL_WIDTH)
                while len(widths) <= c:
                    widths.append(0)
                if sprint and _is_break(text):
                    widths[c] = widths[c] or want
                else:
                    widths[c] = max(widths[c], want)
    return [w or MIN_COL_WIDTH for w in widths]


def text_to_xlsx(text: str, sheet_title: str = "Content") -> bytes:
    """Convert markdown text into an .xlsx workbook and return its bytes."""
    if not text:
        raise ValueError("text must be a non-empty string")

    blocks = _split_blocks(latex_to_text(text))
    has_sprint = any(b["type"] == "table" and _is_sprint_table(b["headers"]) for b in blocks)

    widths = _column_widths(blocks)
    total_cols = len(widths)
    if not widths:
        widths, total_cols = [PROSE_WIDTH], 1
    elif sum(widths) < 60:
        # Narrow tables only: widen the last column so prose rows stay readable.
        widths[-1] += 60 - sum(widths)

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]\*\?/\\:]", "", sheet_title)[:31] or "Content"
    ws.sheet_view.showGridLines = False
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w

    writer = _SheetWriter(ws, total_cols, widths)
    for b in blocks:
        kind = b["type"]
        if kind == "header1":
            writer.prose(_plain(b["text"]), bold=True, size=16, color="111827")
        elif kind == "header2":
            writer.prose(_plain(b["text"]), bold=True, size=14, color="111827")
        elif kind == "header3":
            writer.prose(_plain(b["text"]), bold=True, size=12, color="111827")
        elif kind == "paragraph":
            joined = "\n".join(line.strip() for line in b["lines"])
            writer.prose(_plain(joined), bold=_is_all_bold(joined.strip()))
        elif kind == "ul":
            for i, item in enumerate(b["items"]):
                writer.prose("•  " + _plain(item), gap_after=1 if i == len(b["items"]) - 1 else 0)
        elif kind == "ol":
            for i, item in enumerate(b["items"]):
                writer.prose(f"{i + 1}.  " + _plain(item), gap_after=1 if i == len(b["items"]) - 1 else 0)
        elif kind == "quote":
            writer.prose(_plain("\n".join(b["lines"])), italic=True, color="4B5563")
        elif kind == "code":
            writer.prose(b["text"], mono=True, size=10, fill="F1F5F9")
        elif kind == "image":
            writer.prose(f"[Image: {b.get('alt') or 'figure'}]", italic=True, color="6B7280")
        elif kind == "table":
            sprint = _is_sprint_table(b["headers"])
            # The sprint legend is a one-column table of activity labels: colour it too.
            legend = has_sprint and not sprint and len(b["headers"]) == 1
            writer.table(b["headers"], b["rows"], sprint=sprint, sprint_doc=legend)
        # "rule" blocks are just spacing in a spreadsheet.

    if has_sprint:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
